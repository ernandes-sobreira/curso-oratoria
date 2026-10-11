"""Synthetic test presentation: none of these values are research data."""
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pathlib import Path
import sys
out=Path(sys.argv[1] if len(sys.argv)>1 else '/tmp/oratoria-fixture.pptx')
out.parent.mkdir(parents=True, exist_ok=True)
p=Presentation()
s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='Fogo e água: exemplo sintético'
s.placeholders[1].text='Série inventada exclusivamente para teste: 2024 = 0; 2025 = 12,5.\nFonte: fixture automatizada.\nTexto literal <img src=x onerror=alert(1)> e símbolos & %.'
s.notes_slide.notes_text_frame.text='Notas de teste: não extrapolar nem apresentar como pesquisa real.'
s=p.slides.add_slide(p.slide_layouts[5]);s.shapes.title.text='Água e fogo: gráfico sintético'
data=CategoryChartData();data.categories=['A','B','C'];data.add_series('Fixture sem dados reais',(0,None,12.5));s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(3),data)
s=p.slides.add_slide(p.slide_layouts[5]);s.shapes.title.text='Terceiro slide oculto';s._element.set('show','0')
# Order intentionally differs from slide1.xml, slide2.xml, slide3.xml.
ids=p.slides._sldIdLst; item=ids[1];ids.remove(item);ids.insert(0,item)
p.save(out)
print(out)
